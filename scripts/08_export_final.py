#!/usr/bin/env python3
"""
Export the final gRNA table to both CSV and Excel (.xlsx).
Applies conditional formatting in Excel for easy interpretation.
"""

import pandas as pd
import argparse
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

FINAL_COLUMNS = [
    "rank",
    "grna_id",
    "spacer",           # 20 bp spacer sequence (5′→3′)
    "pam",              # NGG trinucleotide
    "strand",           # + or -
    "position_nt",      # Position in target gene
    "gc_content",       # Spacer GC fraction (0–1)
    "seed_gc",          # Seed region (positions 9–20) GC fraction
    "mfe_kcal",         # RNAfold MFE (kcal/mol)
    "mfe_score",        # MFE-derived score (0–1)
    "ot_0mm_adjusted",  # Off-targets with 0 mismatches (excl. on-target)
    "ot_1mm",           # Off-targets with 1 mismatch
    "ot_2mm",           # Off-targets with 2 mismatches
    "ot_3mm",           # Off-targets with 3 mismatches
    "total_offtargets", # Total Cas-OFFinder off-targets (0–3mm)
    "blast_offtarget_hits",  # Corroborating BLAST off-target hits
    "efficiency_score", # Composite score (0–1; higher = better)
    "verdict",          # TOP CANDIDATE / CANDIDATE / MARGINAL / REJECT
]

FILL_TOP    = PatternFill("solid", fgColor="C6EFCE")   # Green
FILL_CAND   = PatternFill("solid", fgColor="FFEB9C")   # Yellow
FILL_MARGIN = PatternFill("solid", fgColor="FFCC99")   # Orange
FILL_REJECT = PatternFill("solid", fgColor="FFC7CE")   # Red

VERDICT_FILL = {
    "TOP CANDIDATE"           : FILL_TOP,
    "CANDIDATE"               : FILL_CAND,
    "MARGINAL"                : FILL_MARGIN,
}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input",  required=True)
    parser.add_argument("--output", required=True, help="Output prefix (no extension)")
    args = parser.parse_args()

    df = pd.read_csv(args.input)

    # Keep only defined columns that exist in dataframe
    cols = [c for c in FINAL_COLUMNS if c in df.columns]
    df_final = df[cols].copy()

    # Save CSV
    csv_path  = args.output + ".csv"
    xlsx_path = args.output + ".xlsx"
    df_final.to_csv(csv_path, index=False)
    print(f"[OK] CSV saved: {csv_path}")

    # Save Excel with formatting
    df_final.to_excel(xlsx_path, index=False, sheet_name="gRNA_Results")

    wb = load_workbook(xlsx_path)
    ws = wb["gRNA_Results"]

    # Header formatting
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Auto column widths
    for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or ""))
             for r in range(1, ws.max_row + 1)),
            default=10
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 30)

    # Row colouring by verdict
    verdict_col_idx = cols.index("verdict") + 1 if "verdict" in cols else None
    if verdict_col_idx:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            verdict_val = str(row[verdict_col_idx - 1].value or "")
            fill = VERDICT_FILL.get(verdict_val, FILL_REJECT)
            for cell in row:
                cell.fill = fill

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(xlsx_path)
    print(f"[OK] Excel saved with formatting: {xlsx_path}")

    # Summary
    vc = df_final["verdict"].value_counts() if "verdict" in df_final.columns else {}
    print("\n── Final Verdict Summary ──────────────────")
    for v, n in vc.items():
        print(f"  {v:<35} {n:>3}")

if __name__ == "__main__":
    main()
